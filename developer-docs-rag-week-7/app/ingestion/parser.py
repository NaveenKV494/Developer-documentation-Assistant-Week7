from __future__ import annotations

from dataclasses import dataclass
import io
from pathlib import Path
import re
from typing import BinaryIO


@dataclass
class ParsedBlock:
    text: str
    page: int = 1
    section: str = ""
    kind: str = "text"  # "text", "heading", "code", "table"


FENCE_RE = re.compile(r"^\s*(```+|~~~+)\s*(.*)$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


def _read_bytes(source: str | Path | BinaryIO | bytes) -> tuple[bytes, str]:
    """Return raw bytes and a resolved filename if available."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")
        return path.read_bytes(), path.name
    elif isinstance(source, bytes):
        return source, ""
    elif hasattr(source, "read"):
        filename = getattr(source, "name", "")
        data = source.read()
        if hasattr(source, "seek"):
            source.seek(0)
        return data, filename
    raise TypeError(f"Unsupported source type: {type(source)}")


def parse_pdf(source: str | Path | BinaryIO | bytes) -> list[ParsedBlock]:
    """Parse PDF page-by-page using pypdf."""
    from pypdf import PdfReader

    data, _ = _read_bytes(source)
    stream = io.BytesIO(data)
    reader = PdfReader(stream)

    if not reader.pages:
        raise ValueError("PDF contains no pages")

    blocks: list[ParsedBlock] = []
    for page_idx, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        page_text = page_text.strip()
        if not page_text:
            continue

        lines = [line.strip() for line in page_text.splitlines() if line.strip()]
        # Guess section from the first short line of the page if it looks like a title
        first_line = lines[0] if lines else ""
        section_title = first_line if (first_line and len(first_line) < 80 and not first_line.endswith(".")) else f"Page {page_idx}"

        # Group page into paragraphs
        paragraphs = re.split(r"\n\s*\n", page_text)
        current_section = section_title
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            # If paragraph is very short and title-like, it might be an internal heading
            if len(para) < 60 and "\n" not in para and not para.endswith("."):
                current_section = para
                blocks.append(ParsedBlock(text=para, page=page_idx, section=current_section, kind="heading"))
            else:
                blocks.append(ParsedBlock(text=para, page=page_idx, section=current_section, kind="text"))

    if not blocks:
        raise ValueError("No extractable text found in PDF (it may be scanned or empty)")

    return blocks


def parse_docx(source: str | Path | BinaryIO | bytes) -> list[ParsedBlock]:
    """Parse DOCX document preserving headings as sections."""
    import docx

    data, _ = _read_bytes(source)
    doc = docx.Document(io.BytesIO(data))

    blocks: list[ParsedBlock] = []
    current_section = ""

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name.lower() if para.style else ""
        if "heading" in style_name or "title" in style_name:
            current_section = text
            blocks.append(ParsedBlock(text=text, page=1, section=current_section, kind="heading"))
        else:
            blocks.append(ParsedBlock(text=text, page=1, section=current_section, kind="text"))

    # Also extract tables
    for table in doc.tables:
        rows_text = []
        for row in table.rows:
            cell_texts = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows_text.append(" | ".join(cell_texts))
        if rows_text:
            table_content = "\n".join(rows_text)
            blocks.append(ParsedBlock(text=table_content, page=1, section=current_section, kind="table"))

    if not blocks:
        raise ValueError("No extractable text found in DOCX file")

    return blocks


def parse_markdown(source: str | Path | BinaryIO | bytes) -> list[ParsedBlock]:
    """Parse Markdown text into semantic blocks with headings, tables, and code."""
    data, _ = _read_bytes(source)
    text = data.decode("utf-8", errors="replace")

    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[ParsedBlock] = []
    paragraph: list[str] = []
    current_heading = ""
    i = 0

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            val = "\n".join(paragraph).strip()
            if val:
                blocks.append(ParsedBlock(text=val, page=1, section=current_heading, kind="text"))
            paragraph = []

    while i < len(lines):
        line = lines[i]
        fence = FENCE_RE.match(line)
        if fence:
            flush_paragraph()
            fence_marker = fence.group(1)
            fence_lines = [line]
            i += 1
            while i < len(lines):
                fence_lines.append(lines[i])
                if lines[i].lstrip().startswith(fence_marker[0] * len(fence_marker)):
                    i += 1
                    break
                i += 1
            blocks.append(ParsedBlock(
                text="\n".join(fence_lines).strip(),
                page=1,
                section=current_heading,
                kind="code"
            ))
            continue

        heading = HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            current_heading = heading.group(2).strip()
            blocks.append(ParsedBlock(
                text=line.strip(),
                page=1,
                section=current_heading,
                kind="heading"
            ))
            i += 1
            continue

        # Markdown table
        if i + 1 < len(lines) and "|" in line and TABLE_SEPARATOR_RE.match(lines[i + 1]):
            flush_paragraph()
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                table_lines.append(lines[i])
                i += 1
            blocks.append(ParsedBlock(
                text="\n".join(table_lines).strip(),
                page=1,
                section=current_heading,
                kind="table"
            ))
            continue

        if not line.strip():
            flush_paragraph()
            i += 1
            continue

        paragraph.append(line)
        i += 1

    flush_paragraph()

    if not blocks:
        raise ValueError("No content found in Markdown document")

    return blocks


def parse_txt(source: str | Path | BinaryIO | bytes) -> list[ParsedBlock]:
    """Parse plain text into paragraphs."""
    data, _ = _read_bytes(source)
    text = data.decode("utf-8", errors="replace")

    paragraphs = re.split(r"\n\s*\n", text)
    blocks: list[ParsedBlock] = []
    current_section = ""

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        # If very short single line without trailing period, treat as section heading
        if len(para) < 60 and "\n" not in para and not para.endswith("."):
            current_section = para
            blocks.append(ParsedBlock(text=para, page=1, section=current_section, kind="heading"))
        else:
            blocks.append(ParsedBlock(text=para, page=1, section=current_section, kind="text"))

    if not blocks:
        raise ValueError("No content found in TXT document")

    return blocks


def parse_document(source: str | Path | BinaryIO | bytes, filename: str | None = None) -> list[ParsedBlock]:
    """Unified entry point to parse PDF, DOCX, MD, or TXT."""
    if not filename:
        if isinstance(source, (str, Path)):
            filename = Path(source).name
        elif hasattr(source, "name"):
            filename = getattr(source, "name", "")
        else:
            filename = "document.txt"

    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return parse_pdf(source)
    elif ext in (".docx", ".doc"):
        return parse_docx(source)
    elif ext in (".md", ".markdown"):
        return parse_markdown(source)
    elif ext in (".txt", ".text", ""):
        return parse_txt(source)
    else:
        # Fallback to plain text parsing for unknown extensions
        return parse_txt(source)
